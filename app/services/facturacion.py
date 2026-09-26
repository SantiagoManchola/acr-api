"""Procesamiento en memoria del formato mensual de facturación.

El servicio solo consulta la DB. Nunca crea ni modifica suscriptores, medidores
o lecturas. La salida es una copia del XLSX original en la que únicamente se
rellena la columna del mes solicitado.
"""
from collections import Counter, defaultdict
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
import math
import re
import unicodedata
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import models

MESES = (
    ("ene", "Enero"), ("feb", "Febrero"), ("mar", "Marzo"),
    ("abr", "Abril"), ("may", "Mayo"), ("jun", "Junio"),
    ("jul", "Julio"), ("ago", "Agosto"), ("sep", "Septiembre"),
    ("oct", "Octubre"), ("nov", "Noviembre"), ("dic", "Diciembre"),
)
MES_POR_NUMERO = {i + 1: item for i, item in enumerate(MESES)}
MESES_VALIDOS = {abreviatura for abreviatura, _ in MESES}
ENCABEZADOS_REQUERIDOS = {
    "area", "nombre", "nmedidor", "codruta", "mesant", *MESES_VALIDOS,
}
MAX_ARCHIVO_BYTES = 20 * 1024 * 1024
MAX_FILAS = 20_000


class PlantillaFacturacionError(ValueError):
    """Error legible de validación de la plantilla de facturación."""


def _codigo(valor) -> str:
    """Normaliza códigos numéricos de Excel (texto, int o float terminado en .0)."""
    if valor is None or isinstance(valor, bool):
        return ""
    if isinstance(valor, (int, float)) and math.isfinite(float(valor)) and float(valor).is_integer():
        texto = str(int(valor))
    else:
        texto = str(valor).strip()
        if re.fullmatch(r"\d+\.0+", texto):
            texto = texto.split(".", 1)[0]
    if texto.isdigit():
        return texto.lstrip("0") or "0"
    return texto.strip().upper()


def _texto_normalizado(valor) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c)).upper()
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z0-9 ]", " ", texto)).strip()


def _tokens_nombre(valor) -> tuple[str, ...]:
    return tuple(sorted(_texto_normalizado(valor).split()))


def _numero_excel(valor):
    """Decimal finito para una celda numérica o un resultado de fórmula cacheado."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, str):
        valor = valor.strip()
        if not valor or valor.startswith("="):
            return None
        # Los archivos contables pueden usar coma decimal en celdas de texto.
        if "," in valor and "." not in valor:
            valor = valor.replace(",", ".")
    try:
        n = Decimal(str(valor))
    except (InvalidOperation, ValueError, TypeError):
        return None
    return n if n.is_finite() else None


def _valor_json(valor):
    if isinstance(valor, Decimal):
        return float(valor)
    if isinstance(valor, (date,)):
        return valor.isoformat()
    if valor is None or isinstance(valor, (str, int, float, bool)):
        return valor
    if hasattr(valor, "value"):
        return str(valor.value)
    return str(valor)


def _buscar_plantilla(contenido: bytes):
    try:
        wb = load_workbook(BytesIO(contenido), data_only=False)
        wb_cache = load_workbook(BytesIO(contenido), data_only=True)
    except (InvalidFileException, BadZipFile, OSError, ValueError) as exc:
        raise PlantillaFacturacionError(
            "No se pudo abrir el archivo. Sube una plantilla válida de Excel .xlsx."
        ) from exc

    candidatas = []
    for ws in wb.worksheets:
        for fila_encabezado in range(1, min(ws.max_row or 1, 15) + 1):
            headers = {}
            for col in range(1, ws.max_column + 1):
                value = ws.cell(fila_encabezado, col).value
                if value is not None:
                    headers[str(value).strip().lower()] = col
            if ENCABEZADOS_REQUERIDOS.issubset(headers):
                candidatas.append((ws.title, fila_encabezado, headers))
                break

    if not candidatas:
        wb.close()
        wb_cache.close()
        faltan = sorted(ENCABEZADOS_REQUERIDOS)
        raise PlantillaFacturacionError(
            "No encontré una hoja con los encabezados del formato contable "
            f"(se requieren area, nombre, nmedidor, codruta, mesant y meses ene-dic)."
        )
    if len(candidatas) > 1:
        wb.close()
        wb_cache.close()
        nombres = ", ".join(x[0] for x in candidatas)
        raise PlantillaFacturacionError(
            f"Hay varias hojas que parecen el formato de facturación ({nombres}); deja solo una."
        )

    nombre_hoja, fila_encabezado, headers = candidatas[0]
    if (wb[nombre_hoja].max_row or 0) - fila_encabezado > MAX_FILAS:
        wb.close()
        wb_cache.close()
        raise PlantillaFacturacionError(
            f"La plantilla supera el límite de {MAX_FILAS:,} filas de datos."
        )
    return wb, wb_cache, wb[nombre_hoja], wb_cache[nombre_hoja], fila_encabezado, headers


def _lectura_sugerida(db: Session, fecha_desde: date, fecha_hasta: date):
    suscriptores = db.execute(select(models.Suscriptor)).scalars().all()
    medidores = db.execute(select(models.Micromedidor)).scalars().all()

    por_codigo = {_codigo(s.codigo_usuario): s for s in suscriptores if _codigo(s.codigo_usuario)}
    por_facturacion = defaultdict(list)
    for s in suscriptores:
        cod = _codigo(s.codigo_facturacion)
        if cod:
            por_facturacion[cod].append(s)

    medidor_por_serial = {_codigo(m.serial): m for m in medidores if _codigo(m.serial)}
    medidores_por_suscriptor = defaultdict(list)
    for m in medidores:
        if m.suscriptor_id is not None:
            medidores_por_suscriptor[m.suscriptor_id].append(m)

    ids_medidores = [m.id for m in medidores]
    lecturas_periodo = defaultdict(list)
    if ids_medidores:
        stmt = (
            select(models.Lectura)
            .where(
                models.Lectura.micromedidor_id.in_(ids_medidores),
                models.Lectura.fecha >= fecha_desde,
                models.Lectura.fecha <= fecha_hasta,
            )
            .order_by(models.Lectura.fecha.desc(), models.Lectura.id.desc())
        )
        for lectura in db.execute(stmt).scalars().all():
            lecturas_periodo[lectura.micromedidor_id].append(lectura)

    ultima_global = {}
    if ids_medidores:
        stmt = (
            select(models.Lectura)
            .where(models.Lectura.micromedidor_id.in_(ids_medidores))
            .order_by(models.Lectura.fecha.desc(), models.Lectura.id.desc())
        )
        for lectura in db.execute(stmt).scalars().all():
            ultima_global.setdefault(lectura.micromedidor_id, lectura)

    return {
        "suscriptores": suscriptores,
        "por_codigo": por_codigo,
        "por_facturacion": por_facturacion,
        "medidores": medidores,
        "medidor_por_serial": medidor_por_serial,
        "medidores_por_suscriptor": medidores_por_suscriptor,
        "lecturas_periodo": lecturas_periodo,
        "ultima_global": ultima_global,
    }


def analizar_facturacion(
    db: Session,
    contenido: bytes,
    mes: int,
    anio: int,
    fecha_desde: date,
    fecha_hasta: date,
):
    """Previsualiza el archivo y calcula la acción sugerida por cada fila."""
    if mes not in MES_POR_NUMERO:
        raise PlantillaFacturacionError("Selecciona un mes válido entre enero y diciembre.")
    if anio < 2000 or anio > 2100:
        raise PlantillaFacturacionError("El año de facturación no es válido.")
    if fecha_desde > fecha_hasta:
        raise PlantillaFacturacionError("La fecha inicial debe ser anterior o igual a la fecha final.")
    if not contenido:
        raise PlantillaFacturacionError("Selecciona el archivo Excel de facturación.")
    if len(contenido) > MAX_ARCHIVO_BYTES:
        raise PlantillaFacturacionError("El archivo supera el máximo permitido de 20 MB.")

    wb, wb_cache, ws, ws_cache, fila_encabezado, headers = _buscar_plantilla(contenido)
    datos_db = _lectura_sugerida(db, fecha_desde, fecha_hasta)
    col_mes = headers[MES_POR_NUMERO[mes][0]]
    col_mes_prev = headers[MESES[mes - 2][0]] if mes >= 2 else None
    col_mes_prevprev = headers[MESES[mes - 3][0]] if mes >= 3 else None

    conteo_areas = Counter(
        _codigo(ws.cell(r, headers["area"]).value)
        for r in range(fila_encabezado + 1, ws.max_row + 1)
        if _fila_con_datos(ws, r, headers)
    )

    auditoria = []
    novedades = []
    planes = []
    existentes = 0
    total_filas = 0

    def avisar(tipo, fila, area, nombre, nmedidor_excel, detalle, **extra):
        novedades.append({
            "tipo": tipo,
            "fila": fila,
            "area": str(area or ""),
            "nombre_excel": str(nombre or ""),
            "nmedidor_excel": _valor_json(nmedidor_excel),
            "detalle": detalle,
            **{k: _valor_json(v) for k, v in extra.items()},
        })

    for fila in range(fila_encabezado + 1, ws.max_row + 1):
        if not _fila_con_datos(ws, fila, headers):
            continue
        total_filas += 1
        area = ws.cell(fila, headers["area"]).value
        nombre_excel = ws.cell(fila, headers["nombre"]).value
        nmedidor_excel = ws.cell(fila, headers["nmedidor"]).value
        codruta_excel = ws.cell(fila, headers["codruta"]).value
        celda_objetivo = ws.cell(fila, col_mes)
        valor_existente = celda_objetivo.value
        valor_existente_cache = ws_cache.cell(fila, col_mes).value
        # El 0 y las celdas vacías son el "espacio a generar" de la plantilla
        # contable: no cuentan como dato existente ni disparan advertencias.
        valor_existente_num = _numero_excel(valor_existente_cache)
        existe_valor_real = valor_existente_num is not None and valor_existente_num != 0
        if existe_valor_real:
            existentes += 1

        row_issues = []
        def issue(tipo, detalle, **extra):
            row_issues.append(tipo)
            avisar(tipo, fila, area, nombre_excel, nmedidor_excel, detalle, **extra)

        area_key = _codigo(area)
        ruta_key = _codigo(codruta_excel)
        sub = datos_db["por_codigo"].get(area_key) if area_key else None
        if sub is None and ruta_key:
            rutas = datos_db["por_facturacion"].get(ruta_key, [])
            if len(rutas) == 1:
                sub = rutas[0]
                issue(
                    "coincidencia_por_codruta",
                    "area no encontró suscriptor; se identificó por código de facturación codruta.",
                    codigo_db=sub.codigo_usuario,
                )

        if sub is None:
            issue(
                "usuario_no_encontrado",
                "El área/código no está en suscriptores de la DB. Se conserva la fila y se revisa por grupo.",
            )

        elif _tokens_nombre(nombre_excel) != _tokens_nombre(sub.nombre):
            issue(
                "nombre_diferente",
                "El código de cuenta coincide, pero el nombre del Excel difiere del nombre en la DB.",
                nombre_db=sub.nombre,
                codigo_db=sub.codigo_usuario,
            )

        if sub is not None and ruta_key:
            ruta_db = _codigo(sub.codigo_facturacion)
            if ruta_db and ruta_db != ruta_key:
                issue(
                    "codruta_diferente",
                    "codruta del Excel no coincide con codigo_facturacion de la cuenta encontrada por area.",
                    codruta_db=sub.codigo_facturacion,
                    codigo_db=sub.codigo_usuario,
                )

        if sub is not None and area_key and conteo_areas[area_key] > 1:
            issue(
                "area_repetida",
                "El código area aparece más de una vez en la plantilla; revisar que sean filas intencionales.",
                filas_con_area=conteo_areas[area_key],
            )

        if sub is not None and sub.estado == models.EstadoRegistro.inactivo:
            issue(
                "suscriptor_inactivo",
                "El suscriptor está marcado inactivo en la DB; se conserva porque aparece en el formato contable.",
                codigo_db=sub.codigo_usuario,
            )

        valor_destino = None
        origen = "pendiente"
        serial_usado = None
        fecha_lectura = None
        valor_sugerido = None
        usuario_id = sub.id if sub is not None else None
        candidatos = datos_db["medidores_por_suscriptor"].get(usuario_id, []) if usuario_id else []
        activos = [m for m in candidatos if m.estado == models.EstadoRegistro.activo]
        serial_key = _codigo(nmedidor_excel)
        medidor_exact = datos_db["medidor_por_serial"].get(serial_key) if serial_key else None
        medidor = None

        if sub is not None:
            if medidor_exact and medidor_exact.suscriptor_id == sub.id:
                medidor = medidor_exact
            else:
                if serial_key:
                    if medidor_exact and medidor_exact.suscriptor_id != sub.id:
                        issue(
                            "serial_otro_suscriptor",
                            "nmedidor del Excel está registrado a nombre de otro suscriptor; se buscará el medidor asociado a esta cuenta.",
                            suscriptor_serial_db=medidor_exact.suscriptor_id,
                        )
                    else:
                        issue(
                            "serial_no_en_db",
                            "nmedidor del Excel no coincide con un serial registrado; se buscará el medidor asociado a esta cuenta.",
                        )

                if len(activos) == 1:
                    medidor = activos[0]
                    if not serial_key:
                        issue(
                            "nmedidor_vacio_excel",
                            "Excel no trae nmedidor, pero la DB tiene un único medidor activo asociado; se usará ese medidor.",
                            serial_db=medidor.serial,
                            codigo_db=sub.codigo_usuario,
                        )
                    elif _codigo(medidor.serial) != serial_key:
                        issue(
                            "serial_diferente_db",
                            "Se usará el único medidor activo asociado por la DB; el serial difiere del Excel.",
                            serial_db=medidor.serial,
                            codigo_db=sub.codigo_usuario,
                        )
                elif not activos and not candidatos:
                    if serial_key:
                        issue(
                            "medidor_excel_sin_asociacion",
                            "La fila indica un medidor, pero la cuenta no tiene medidor asociado en la DB; no se adivina una lectura.",
                            codigo_db=sub.codigo_usuario,
                        )
                    else:
                        # Sin medidor en la DB ni en el formato: estimar con el historial
                        # de la propia plantilla. Si el usuario no está en la DB, se aplica
                        # la misma regla y se conserva la novedad usuario_no_encontrado.
                        prev = _numero_excel(ws_cache.cell(fila, col_mes_prev).value) if col_mes_prev else None
                        prevprev = _numero_excel(ws_cache.cell(fila, col_mes_prevprev).value) if col_mes_prevprev else None
                        if mes >= 3 and prev is not None and prevprev is not None:
                            valor_destino = float(prev - prevprev + prev)
                            valor_sugerido = valor_destino
                            origen = "formula_sin_medidor"
                        else:
                            issue(
                                "historial_insuficiente_formula",
                                "No se puede calcular la fórmula: se requieren dos meses previos con valores numéricos (disponible desde marzo).",
                            )
                elif len(activos) > 1:
                    issue(
                        "multiples_medidores",
                        "La cuenta tiene más de un medidor activo y el nmedidor del Excel no identifica uno de ellos; requiere selección manual.",
                        seriales_db=", ".join(str(m.serial) for m in activos),
                        codigo_db=sub.codigo_usuario,
                    )
                elif len(candidatos) == 1:
                    # Solo hay un medidor asociado, aunque esté inactivo; puede tener
                    # una lectura válida del periodo que se está facturando.
                    medidor = candidatos[0]
                    issue(
                        "medidor_inactivo_db",
                        "La cuenta tiene un único medidor asociado, pero está inactivo en la DB; se revisará si existe lectura en el periodo.",
                        serial_db=medidor.serial,
                        codigo_db=sub.codigo_usuario,
                    )
                else:
                    issue(
                        "multiples_medidores_inactivos",
                        "La cuenta tiene varios medidores asociados y el Excel no permite identificar cuál corresponde a este periodo.",
                        seriales_db=", ".join(str(m.serial) for m in candidatos),
                    )

        elif not serial_key:
            # Usuario no encontrado en DB y Excel sin medidor: la regla de estimación
            # depende exclusivamente de los meses previos del propio formato.
            prev = _numero_excel(ws_cache.cell(fila, col_mes_prev).value) if col_mes_prev else None
            prevprev = _numero_excel(ws_cache.cell(fila, col_mes_prevprev).value) if col_mes_prevprev else None
            if mes >= 3 and prev is not None and prevprev is not None:
                valor_destino = float(prev - prevprev + prev)
                valor_sugerido = valor_destino
                origen = "formula_sin_medidor"
            else:
                issue(
                    "historial_insuficiente_formula",
                    "Usuario no encontrado en DB y no hay dos meses previos numéricos para aplicar la fórmula.",
                )

        if medidor is not None:
            serial_usado = str(medidor.serial)
            if medidor.estado == models.EstadoRegistro.inactivo:
                issue(
                    "medidor_inactivo_db",
                    "El medidor seleccionado está inactivo en la DB; se usará solo si tiene lectura dentro del rango.",
                    serial_db=serial_usado,
                    codigo_db=sub.codigo_usuario if sub else None,
                )
            lecturas = datos_db["lecturas_periodo"].get(medidor.id, [])
            if lecturas:
                lectura = lecturas[0]  # ya está ordenado por fecha DESC, id DESC
                valor_destino = float(lectura.lectura)
                valor_sugerido = valor_destino
                fecha_lectura = lectura.fecha
                origen = "lectura_db"
                if len(lecturas) > 1:
                    issue(
                        "varias_lecturas_periodo",
                        f"Hay {len(lecturas)} lecturas en el rango; se tomó la más reciente ({lectura.fecha}).",
                        serial_db=serial_usado,
                        fecha_lectura=lectura.fecha,
                        valor_sugerido=valor_destino,
                    )
                if lectura.promedio_usado:
                    issue(
                        "lectura_estimada_db",
                        "La lectura seleccionada en la DB fue estimada mediante promedio.",
                        serial_db=serial_usado,
                        fecha_lectura=lectura.fecha,
                        valor_sugerido=valor_destino,
                    )
                if lectura.irregular:
                    issue(
                        "lectura_irregular_db",
                        "La lectura seleccionada está marcada como irregular en la DB.",
                        serial_db=serial_usado,
                        fecha_lectura=lectura.fecha,
                        valor_sugerido=valor_destino,
                    )
            else:
                ultima = datos_db["ultima_global"].get(medidor.id)
                issue(
                    "sin_lectura_en_periodo",
                    "El suscriptor tiene medidor, pero no hay lectura dentro del rango elegido.",
                    serial_db=serial_usado,
                    ultima_fecha_db=ultima.fecha if ultima else None,
                    ultima_lectura_db=ultima.lectura if ultima else None,
                )

        if valor_destino is not None and existe_valor_real:
            anterior_num = _numero_excel(valor_existente_cache)
            if anterior_num != Decimal(str(valor_destino)):
                issue(
                    "valor_mes_difiere",
                    "El Excel ya traía un valor distinto de cero en este mes; en la copia se reemplazará por el valor ACR.",
                    valor_excel=anterior_num,
                    valor_sugerido=valor_destino,
                    serial_db=serial_usado,
                    fecha_lectura=fecha_lectura,
                )

        if valor_destino is None and not row_issues:
            issue("sin_resultado", "No se pudo determinar un valor para esta fila.")

        planes.append({
            "fila": fila,
            "columna_objetivo": col_mes,
            "valor_original": valor_existente,
            "valor_destino": valor_destino,
            "origen": origen,
        })
        auditoria.append({
            "fila": fila,
            "area": str(area or ""),
            "nombre_excel": str(nombre_excel or ""),
            "codigo_db": str(sub.codigo_usuario) if sub else None,
            "nombre_db": str(sub.nombre) if sub else None,
            "nmedidor_excel": _valor_json(nmedidor_excel),
            "serial_usado_db": serial_usado,
            "origen_valor": origen,
            "fecha_lectura_db": _valor_json(fecha_lectura),
            "valor_excel_original": _valor_json(valor_existente_cache),
            "valor_generado": _valor_json(valor_destino),
            "estado_fila": "completada" if valor_destino is not None else "pendiente",
            "novedades": "; ".join(row_issues),
        })

    resumen = {
        "hoja": ws.title,
        "fila_encabezado": fila_encabezado,
        "mes": mes,
        "mes_nombre": MES_POR_NUMERO[mes][1],
        "anio": anio,
        "fecha_desde": fecha_desde.isoformat(),
        "fecha_hasta": fecha_hasta.isoformat(),
        "filas_total": total_filas,
        "filas_con_lectura_db": sum(1 for p in planes if p["origen"] == "lectura_db"),
        "filas_formula": sum(1 for p in planes if p["origen"] == "formula_sin_medidor"),
        "filas_pendientes": sum(1 for p in planes if p["valor_destino"] is None),
        "usuarios_no_encontrados": sum(1 for n in novedades if n["tipo"] == "usuario_no_encontrado"),
        "celdas_mes_con_datos": existentes,
        "novedades_total": len(novedades),
    }
    wb.close()
    wb_cache.close()
    return {"resumen": resumen, "novedades": novedades, "auditoria": auditoria, "planes": planes}


def generar_archivo_facturacion(contenido: bytes, hoja: str, planes: list[dict]) -> bytes:
    """Rellena solo la columna objetivo y devuelve un XLSX nuevo en memoria."""
    try:
        wb = load_workbook(BytesIO(contenido), data_only=False)
    except (InvalidFileException, BadZipFile, OSError, ValueError) as exc:
        raise PlantillaFacturacionError("No se pudo volver a abrir la plantilla para generar la salida.") from exc
    if hoja not in wb.sheetnames:
        wb.close()
        raise PlantillaFacturacionError("La hoja de la plantilla cambió entre la previsualización y la descarga.")
    ws = wb[hoja]
    for p in planes:
        celda = ws.cell(p["fila"], p["columna_objetivo"])
        # El output es una copia nueva. Los casos sin resultado quedan vacíos
        # para evitar importar un dato antiguo no verificado como si fuera nuevo.
        celda.value = p["valor_destino"]

    # Mantener cálculo automático para las demás fórmulas que traiga la plantilla.
    try:
        wb.calculation.calcMode = "auto"
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except AttributeError:
        pass
    out = BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


def _fila_con_datos(ws, fila: int, headers: dict) -> bool:
    # Fila de datos solo si contiene algún valor en sus columnas de identificación.
    for key in ("area", "nombre", "nmedidor", "matricula", "codruta"):
        col = headers.get(key)
        if col is None:
            continue
        v = ws.cell(fila, col).value
        if v is not None and str(v).strip():
            return True
    return False
